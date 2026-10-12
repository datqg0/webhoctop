import io
import json
import logging
from typing import List, Dict, Any
from sqlalchemy.orm import Session
from app.models import KnowledgeDocument, KnowledgeChunk
from app.services.gemini_service import gemini_service

logger = logging.getLogger(__name__)

class IngestionService:
    def process_pdf(self, db: Session, title: str, subject_code: str, file_bytes: bytes, file_name: str) -> KnowledgeDocument:
        """
        Parse and index textbook PDF document
        """
        doc = KnowledgeDocument(
            title=title,
            doc_type="pdf",
            subject_code=subject_code,
            file_name=file_name,
            total_chunks=0
        )
        db.add(doc)
        db.flush()

        chunks_created = 0
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                text = text.strip()
                if not text:
                    continue

                # Sliding window chunking (approx 500 characters / chunk)
                words = text.split()
                chunk_size = 120
                overlap = 25
                
                for i in range(0, len(words), chunk_size - overlap):
                    chunk_words = words[i:i + chunk_size]
                    chunk_text = " ".join(chunk_words)
                    if len(chunk_text) < 40:
                        continue
                    
                    emb = gemini_service.get_embedding(chunk_text, task_type="retrieval_document")
                    chunk = KnowledgeChunk(
                        doc_id=doc.id,
                        chunk_text=chunk_text,
                        page_or_slide=f"Trang {page_idx + 1}",
                        section_title=f"Mục trang {page_idx + 1}",
                        embedding_json=json.dumps(emb)
                    )
                    db.add(chunk)
                    chunks_created += 1
        except Exception as e:
            db.rollback()
            logger.error(f"Error parsing PDF: {e}")
            raise ValueError("Kh?ng ??c ???c t?p PDF. H?y ki?m tra t?p v? th? l?i.") from e

        if chunks_created == 0:
            db.rollback()
            raise ValueError("PDF kh?ng c? v?n b?n tr?ch xu?t ???c; t?p scan c?n OCR tr??c khi t?i l?n.")

        doc.total_chunks = chunks_created
        db.commit()
        db.refresh(doc)
        return doc

    def process_pptx(self, db: Session, title: str, subject_code: str, file_bytes: bytes, file_name: str) -> KnowledgeDocument:
        """
        Parse and index presentation slides (PPTX)
        """
        doc = KnowledgeDocument(
            title=title,
            doc_type="slide",
            subject_code=subject_code,
            file_name=file_name,
            total_chunks=0
        )
        db.add(doc)
        db.flush()

        chunks_created = 0
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            for idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                
                full_slide_text = "\n".join(slide_texts)
                if not full_slide_text:
                    continue

                emb = gemini_service.get_embedding(full_slide_text, task_type="retrieval_document")
                chunk = KnowledgeChunk(
                    doc_id=doc.id,
                    chunk_text=full_slide_text,
                    page_or_slide=f"Slide {idx + 1}",
                    section_title=slide_texts[0][:50] if slide_texts else f"Slide {idx + 1}",
                    embedding_json=json.dumps(emb)
                )
                db.add(chunk)
                chunks_created += 1
        except Exception as e:
            db.rollback()
            logger.error(f"Error parsing PPTX: {e}")
            raise ValueError("Kh?ng ??c ???c t?p PPTX. H?y ki?m tra t?p v? th? l?i.") from e

        if chunks_created == 0:
            db.rollback()
            raise ValueError("Slide kh?ng c? v?n b?n tr?ch xu?t ???c.")

        doc.total_chunks = chunks_created
        db.commit()
        db.refresh(doc)
        return doc

    def process_qa_bank(self, db: Session, title: str, subject_code: str, qa_items: List[Dict[str, Any]], file_name: str) -> KnowledgeDocument:
        """
        Index a structured Q&A / Detailed solutions bank
        qa_items format: [{"question": "...", "correct": "...", "explanation": "...", "chapter": "..."}]
        """
        doc = KnowledgeDocument(
            title=title,
            doc_type="qa_bank",
            subject_code=subject_code,
            file_name=file_name,
            total_chunks=0
        )
        db.add(doc)
        db.flush()

        chunks_created = 0
        for idx, item in enumerate(qa_items):
            q_text = item.get("question", "")
            ans = item.get("correct", "")
            exp = item.get("explanation", "")
            chapter = item.get("chapter", "Tổng hợp")

            content = f"Câu hỏi: {q_text}\nĐáp án: {ans}\nGiải thích chi tiết: {exp}"
            emb = gemini_service.get_embedding(content, task_type="retrieval_document")
            
            chunk = KnowledgeChunk(
                doc_id=doc.id,
                chunk_text=content,
                page_or_slide=f"Câu {idx + 1}",
                section_title=f"Chủ đề: {chapter}",
                embedding_json=json.dumps(emb)
            )
            db.add(chunk)
            chunks_created += 1

        doc.total_chunks = chunks_created
        db.commit()
        db.refresh(doc)
        return doc

ingestion_service = IngestionService()
